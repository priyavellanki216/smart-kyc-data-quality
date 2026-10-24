from datetime import date

from openpyxl import load_workbook

from kyc_quality.reporting import build_report_frames, export_reports, summary_metrics
from kyc_quality.service import run_demo_pipeline


def test_exports_create_populated_csv_and_xlsx_reporting_outputs(tmp_path):
    database = tmp_path / "demo.db"
    run_demo_pipeline(database, as_of=date(2026, 10, 24))
    frames = build_report_frames(database)
    required = {"client_quality_summary", "exceptions", "duplicate_candidates", "review_queue", "case_aging", "reviewer_workload", "audit_history", "dim_client", "fact_exceptions", "fact_review_cases"}
    assert required.issubset(frames)
    assert len(frames["dim_client"]) == 40
    assert not frames["fact_exceptions"].empty
    output_dir = tmp_path / "exports"
    paths = export_reports(database, output_dir)
    assert "xlsx_workbook" in paths
    assert all(path.exists() for path in paths.values())
    assert paths["exceptions"].read_text(encoding="utf-8-sig").startswith("exception_id,")
    workbook = load_workbook(paths["xlsx_workbook"], read_only=True)
    assert "exceptions" in workbook.sheetnames
    assert "dim_client" in workbook.sheetnames
    assert workbook["dim_client"].max_row == 41


def test_summary_metrics_are_computed_from_persisted_demo_state(tmp_path):
    database = tmp_path / "demo.db"
    run_demo_pipeline(database, as_of=date(2026, 10, 24))
    metrics = summary_metrics(database, as_of=date(2026, 10, 24))
    assert metrics["total_clients"] == 40
    assert metrics["complete_clients"] + metrics["incomplete_clients"] == 40
    assert metrics["open_exceptions"] > 0
    assert metrics["duplicate_candidates"] == 3
    assert metrics["open_review_cases"] > 0
    assert metrics["overdue_periodic_reviews"] == 1
    assert metrics["aged_open_cases_14d"] >= 0
    assert sum(metrics["exceptions_by_category"].values()) == metrics["open_exceptions"]
